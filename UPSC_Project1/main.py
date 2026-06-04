import os
import shutil
import json
import threading
import zipfile
from typing import List
import pandas as pd
from fastapi import FastAPI, HTTPException, Request, BackgroundTasks, Form, UploadFile, File
from fastapi.responses import HTMLResponse, FileResponse
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from google import genai
from google.genai import types
from pydantic import BaseModel

app = FastAPI()

# Cloud Infrastructure Configurations
SUPABASE_URL = "https://nkkprmkdnxcsstczttmz.supabase.co"
SUPABASE_KEY = "sb_publishable_ehPr5Yj3TQuFpFRJshIbmQ_lnIGzsbS"  # Replace if needed
MY_GEMINI_API_KEY = "AIzaSyDEMMScWI1e-Gd8qhILWGaTPif_dlUGWqw"   # Replace if needed

# Initialize AI Client only
ai_client = genai.Client(api_key=MY_GEMINI_API_KEY)
excel_lock = threading.Lock()
CHUNK_DIR = "upload_chunks"

os.makedirs(CHUNK_DIR, exist_ok=True)

class ProcessFragmentsRequest(BaseModel):
    filenames: List[str]

# Global collection to hold questions in memory since we are skipping the database login completely
global_questions_cache = []

def process_document_with_ai(file_bytes: bytes, mime_type: str):
    try:
        prompt = """
        You are an expert UPSC exam coordinator. Analyze the attached document.
        Extract and list every single independent exam question found within the text or images.
        Translate any Hindi question parts into clean English.
        
        For each question, perform a structured classification:
        1. Determine the Main Topic header (e.g., History, Economics, Current Affairs, Polity, Geography).
        2. Determine a logical Sub-topic header that matches the contextual focus (e.g., Modern Indian History, Macroeconomics, International Relations, Indian Constitution).
        
        Output your response STRICTLY as a valid JSON array of objects. Do not wrap it in markdown block quotes.
        Format layout template: [{"Question": "text", "Main_Topic": "text", "Sub_Topic": "text"}]
        """

        json_schema = {
            "type": "ARRAY",
            "items": {
                "type": "OBJECT",
                "properties": {
                    "Question": {"type": "STRING"},
                    "Main_Topic": {"type": "STRING"},
                    "Sub_Topic": {"type": "STRING"}
                },
                "required": ["Question", "Main_Topic", "Sub_Topic"]
            }
        }

        try:
            response = ai_client.models.generate_content(
                model='gemini-2.5-flash',
                contents=[types.Part.from_bytes(data=file_bytes, mime_type=mime_type), prompt],
                config=types.GenerateContentConfig(
                    response_mime_type="application/json", 
                    response_schema=json_schema,
                    temperature=0.1
                )
            )
            return json.loads(response.text)
        except Exception:
            response = ai_client.models.generate_content(
                model='gemini-2.5-flash',
                contents=[types.Part.from_bytes(data=file_bytes, mime_type=mime_type), prompt]
            )
            clean_text = response.text.replace("```json", "").replace("```", "").strip()
            return json.loads(clean_text)
            
    except Exception as e:
        print(f"❌ Gemini Cloud AI processing exception: {e}")
        return []

def build_topic_separated_zip(items: list) -> str:
    zip_filename = "UPSC_Study_Guides.zip"
    temp_dir = "temp_output"
    os.makedirs(temp_dir, exist_ok=True)
    
    font_main = Font(name="Segoe UI", size=13, bold=True, color="1E3A8A")
    font_sub = Font(name="Segoe UI", size=11, bold=True, italic=True, color="374151")
    font_q = Font(name="Segoe UI", size=11)
    fill_main = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid")
    
    df = pd.DataFrame(items)
    if df.empty:
        # Create a blank fallback sheet if no questions were extracted
        df = pd.DataFrame([{"Main_Topic": "GENERAL OUTLINE", "Sub_Topic": "Miscellaneous", "Question": "No questions extracted yet."}])
        
    df['Main_Topic'] = df['Main_Topic'].astype(str).str.strip().str.upper()
    df['Main_Topic'] = df['Main_Topic'].str.replace(r'[\\/*?:\[\]]', '_', regex=True)
    df['Sub_Topic'] = df['Sub_Topic'].astype(str).str.strip().str.title()
    df.drop_duplicates(subset=['Question'], inplace=True)
    
    generated_files = []
    grouped_main = df.groupby('Main_Topic')
    
    for main_name, main_group in grouped_main:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Study Guide"
        
        ws.column_dimensions['A'].width = 32
        ws.column_dimensions['B'].width = 85
        ws.append(["Topics & Subheadings", "Questions"])
        ws.row_dimensions.font = Font(name="Segoe UI", size=11, bold=True)
        
        ws.append([f"─── {main_name} ───", ""])
        current_row = ws.max_row
        ws.cell(row=current_row, column=1).font = font_main
        ws.cell(row=current_row, column=1).fill = fill_main
        
        grouped_sub = main_group.groupby('Sub_Topic')
        for sub_name, sub_group in grouped_sub:
            ws.append([f"  🔹 {sub_name}", ""])
            ws.cell(row=ws.max_row, column=1).font = font_sub
            
            for idx, q_row in enumerate(sub_group['Question'], start=1):
                ws.append(["", f"{idx}. {q_row}"])
                q_row_idx = ws.max_row
                ws.cell(row=q_row_idx, column=2).font = font_q
                ws.cell(row=q_row_idx, column=2).alignment = Alignment(wrap_text=True)
        
        sanitized_filename = main_name.replace(" ", "_")
        file_path = os.path.join(temp_dir, f"{sanitized_filename}_Study_Guide.xlsx")
        wb.save(file_path)
        generated_files.append(file_path)
        
    with zipfile.ZipFile(zip_filename, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for f_path in generated_files:
            zipf.write(f_path, os.path.basename(f_path))
            os.remove(f_path)
            
    try:
        os.rmdir(temp_dir)
    except Exception:
        pass
        
    return zip_filename

def background_fragment_processing_task(filenames: List[str]):
    global global_questions_cache
    try:
        print(f"⏳ Assembling fragments and starting AI processing...")
        for name in filenames:
            ext = os.path.splitext(name)[-1].lower()
            mime_map = {
                '.pdf': 'application/pdf', 
                '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                '.png': 'image/png', 
                '.jpg': 'image/jpeg', 
                '.jpeg': 'image/jpeg'
            }
            if ext not in mime_map: 
                continue
                
            full_file_path = os.path.join(CHUNK_DIR, f"rebuilt_{name}")
            if os.path.exists(full_file_path):
                with open(full_file_path, "rb") as f:
                    file_content = f.read()
                
                res = process_document_with_ai(file_content, mime_map[ext])
                if res: 
                    with excel_lock:
                        for q in res:
                            global_questions_cache.append({
                                "Main_Topic": q.get('Main_Topic', 'GENERAL OUTLINE'),
                                "Sub_Topic": q.get('Sub_Topic', 'Miscellaneous'),
                                "Question": q.get('Question', '')
                            })
                os.remove(full_file_path)
        print(f"✅ AI processing complete. Total questions ready: {len(global_questions_cache)}")
    except Exception as async_err:
        print(f"❌ Background pipeline error: {async_err}")

@app.get("/", response_class=HTMLResponse)
async def serve_ui():
    with open("index.html", "r", encoding="utf-8") as f:
        return f.read()

@app.post("/upload-chunk")
async def handle_chunk_upload(
    chunk: UploadFile = File(...), 
    filename: str = Form(...), 
    upload_id: str = Form(...), 
    chunk_index: int = Form(...), 
    total_chunks: int = Form(...)
):
    chunk_file_path = os.path.join(CHUNK_DIR, f"{upload_id}_{chunk_index}")
    with open(chunk_file_path, "wb") as f:
        shutil.copyfileobj(chunk.file, f)
        
    if chunk_index == total_chunks - 1:
        final_file_path = os.path.join(CHUNK_DIR, f"rebuilt_{filename}")
        with open(final_file_path, "wb") as master_file:
            for i in range(total_chunks):
                part_path = os.path.join(CHUNK_DIR, f"{upload_id}_{i}")
                with open(part_path, "rb") as part_file:
                    master_file.write(part_file.read())
                os.remove(part_path)
                
    return {"status": "chunk_saved"}

@app.post("/process-fragments")
async def process_fragments_trigger(background_tasks: BackgroundTasks, payload: ProcessFragmentsRequest):
    background_tasks.add_task(background_fragment_processing_task, payload.filenames)
    return {"message": "Processing started in background."}

@app.get("/download")
async def download_personal_excel():
    global global_questions_cache
    if not global_questions_cache:
        raise HTTPException(status_code=400, detail="No processed data found yet. Please wait for the AI to complete or upload documents.")
    
    zip_path = build_topic_separated_zip(global_questions_cache)
    return FileResponse(path=zip_path, filename="UPSC_Topic_Wise_Study_Guides.zip", media_type="application/zip")
